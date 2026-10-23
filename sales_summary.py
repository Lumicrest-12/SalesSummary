import argparse
import sys
from pathlib import Path

import pandas as pd
from openpyxl.chart import BarChart,LineChart,Reference
from openpyxl.styles import Alignment,Font,PatternFill
from openpyxl.utils import get_column_letter

ALIASES = {
    "data":["วันที่","วัน","วันที่ขาย","date","order date","sele date"],
    "product":["สินค้า","ชื่อสินค้า","รายการ","รายการสินค้า","product","item","product name"],
    "qty":["จำนวน","qty","quantity","units"],
    "price":["ราคา","ราคาต่อหน่วย","ราคา/หน่วย","price","unit price"],
    "total":["ยอดรวม","ยอดขาย","รวม","จำนวนเงิน","total","amount","revenue","seles"]
}

FONT = "Arial"
HEADER_FILL = PatternFill("solid",start_color="1F4E78")
HEADER_FONT = Font(name=FONT,bold=True,color="FFFFFF")
BODY_FONT = Font(name=FONT)
MONEY_FMT = "#,##0.00"

#-----------------------------READ FILE---------------------------------------------------
def read_file(path: Path) -> pd.DataFrame:
    ext = path.suffix.lower()
    if ext in (".xlsx",".xlsm",".xls"):
        return pd.read_excel(path)
    if ext in (".csv",".txt"):
        for enc in ("utf-8-sig","cp874","utf-8"):
            try:
                return pd.read_csv(path,encoding=enc)
            except UnicodeDecodeError:
                continue
        raise ValueError("Cann't read csv")
    raise ValueError(f"ไม่รองรับไฟล์นามสกุล {ext} (ใช้ .xlsx, .xls หรือ .csv)")

def map_columns(df:pd.DataFrame) -> pd.DataFrame:
    lookup = []
    for std,names in ALIASES.items():
        for n in names:
            lookup[n.strip().lower()] = std
    rename = {}
    for col in df.columns:
        std = lookup.get(str(col).strip().lower())
        if std and std not in rename.values():
            rename[col] = std
    df = df.rename(columns=rename)

    missing = [c for c in ("date","product") if c not in df.columns]
    if "total" not in df.columns and not {"qty","price"} <= set(df.columns):
        missing.append("ยอดรวม (หรือ จำนวน + ราคา)")
    if missing:
        raise ValueError(
            "ไม่พบคอลั่มที่จำเป็น" + ", ".join(missing)
            + "\n คอลั่มในไฟล์คือ:" +", ".join(map(str,df.columns))

        )
    return df

def clean (df:pd.DataFrame) -> pd.DataFrame:
    df = map_columns(df)
    df["date"] = pd.to_datetime(df["date"],errors="coerce", dayfirst=True)
    df.loc[df["date"].dt.year > 2400,"date"] = df["date"].apply(
        lambda d: d.replace(year = d.year - 543) if pd.notna(d) and d.year >2400 else d
    )
    for col in ("qty","prince","total"):
        if col in df.columns:
            df[col] = pd.to_numeric(
                df[col].astype(str).str.replace(",","", regex = False), errors="coerce"
            )
    if "qty" not in df.columns:
        df["qty"] = 1
    if "total" not in df.columns:
        df["total"] = df["qty"] * df["price"]

    before = len(df)
    df = df.dropna(subset=["date","product","total"])
    df["product"] = df["product"].astype(str).str.strip()
    df.attrs["dropped"] = before - len(df)
    return df


#--------------------------------SUMMARY-----------------------------------------
def summarize(df:pd.DataFrame,top_n: int):
    daily = (
        df.groupby(df["date"].dt.date)
        .agg(ยอดขาย=("total","sum"),จำนวนชิ้น=("qty","sum"),จำนวนรายการ=("total","size"))
        .reset_index()
        .rename(columns={"date":"วันที่"})
    )
    products = (
        df.groupby("product")
        .agg(จำนวนที่ขาย=("qty","sum"),ยอดขาp=("total","sum"))
        .sort_values("ยอดขาย",ascending=False)
        .reset_index()
        .rename(columns={"product":"สินค้า"})
    )
    products.insert(0,"อันดับ",range(1,len(products)+1))
    products["สัดส่วน"] = products["ยอดขาย"]/products["ยอดขาย"].sum()

    best_day = daily.loc[daily["ยอดขาย"].index()]
    overview = pd.DataFrame([
            ("ช่วงข้อมูล", f"{df['date'].min():%d/%m/%Y} - {df['date'].max():%d/%m/%Y}"),
            ("ยอดขายรวม (บาท)", df["total"].sum()),
            ("จำนวนชิ้นที่ขายรวม", df["qty"].sum()),
            ("จำนวนรายการขาย", len(df)),
            ("ยอดขายเฉลี่ยต่อวัน (บาท)", daily["ยอดขาย"].mean()),
            ("วันที่ขายดีที่สุด", f"{best_day['วันที่']:%d/%m/%Y} ({best_day['ยอดขาย']:,.2f} บาท)"),
            ("สินค้าขายดีอันดับ 1", products.iloc[0]["สินค้า"]),
    ],
    columns=["หัวข้อ","คำ"],
    )
    return overview,products,daily

#------------------------------------------WRITE EXEL----------------------------------------------------------------  
def style_sheet(ws,money_cols=(),pct_cols=(),widths=None):
    for cell in ws[1]:
        cell.fill,cell.font = HEADER_FILL,HEADER_FONT
        cell.aligment = Alignment(horizontal="center")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = BODY_FONT
    for idx in money_cols:
        for row in ws.iter_rows(min_row=2,min_col=idk,max_col=idk):
            row[0].number_format = MONEY_FMT
    for idk in pct_cols:
        for row in ws.iter_rows(min_row=2,min_col=idk,max_col=idk):
            row[0].number_format = "0.0%"
    for i,col in enumerate(ws.columns,1):
        longest = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.columns_dimeansions[get_column_letter(i)].width = min(max(longest + 4,12),50)
    ws.freeze_panes = "A2"

def save_excel(overview,products,daily,top_n,out_path:Path):
    with pd.ExcelWriter(out_path,engine="openpyxl") as xw:
        overview.to_excel(xw,sheet_name="สรุปภาพรวม",index=False)
        products.to_excel(xw,sheet_name="สินค้าขายดี",index=False)
        daily.to_excel(xw,sheet_name="ยอดขายรายวัน",index=False)
        wb = xw.book

        ws = wb["สรุปภาพรวม"]
        style_sheet(ws)
        for r in range(2, ws.max_row + 1):
            if isinstance(ws.cell(r,2).value,(int,float)):
                ws.cell(r,2).number_format = MONEY_FMT if "บาท" in ws.cell(r,1).value else "#,##0"

            ws.cell(r,2).alignment = Alignment(horizontal="right")

        ws = wb["สินค้าขายดี"]
        style_sheet(ws,money_cols=(4,), pct_cols=(5,))
        n = min(top_n, len(products))
        chart = BarChart()
        chart.type,chart.title = "bar",f"สินค้าขายดี {n} อันดับแรก (บาท)"
        chart.add_data(Reference(ws,min_col=4, min_row=1, max_row=n+1),titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=2, min_row=2, max_row=n+1))
        chart.legend,chart.height,chart.width = None,9,10
        chart.y_axis.delete = chart.x_axis.delete = False
        ws.add_chart(chart,"H2")

        ws = wb["ยอดขายรายวัน"]
        style_sheet(ws, money_cols=(2,1))
        for r in range(2, ws.max_row + 1):
            ws.cell(r,1).number_format = "DD/MM/YYYY"
        line = LineChart()
        line.title = "ยอดขายรายวัน (บาท)"
        line.add_data(Reference(ws ,min_col=2, min_row=1, max_row=ws.max_row),titles_from_data=True)
        line.set_categories(Reference(ws, min_col=1, min_row=2, max_row=ws.max_row))
        line.legend,line.height,line.width = None,9,10
        line.y_axis.delete = line.x_axis.delete = False
        ws.add_chart(line,"F2")


#--------------------------------------------PROGRAM---------------------------------------------------------------------
def pick_file_dialog() -> str | None:
    try:
        import tkinter as tk
        from tkinter import filedialog

        root = tk.Tk()
        root.withdraw()
        root.attributes("-topmost", True)  # ให้หน้าต่างอยู่บนสุด ไม่ไปซ่อนหลังหน้าต่างอื่น
        root.update()
        path = filedialog.askopenfilename(
            parent=root,
            title="เลือกไฟล์ยอดขาย",
            filetypes=[("Excel / CSV", "*.xlsx *.xls *.csv"), ("ทุกไฟล์", "*.*")],
        )
        root.destroy()
        return path or None
    except Exception as e:
        print("เปิดหน้าต่างเลือกไฟล์ไม่ได้ (" + str(e) + ")")
        print("วิธีแก้: พิมพ์ชื่อไฟล์ต่อท้ายคำสั่ง เช่น  python sales_summary.py ยอดขาย.xlsx")
        print("หรือลากไฟล์ยอดขายมาวางทับไฟล์ sales_summary.py")
        return None


def show_message(title: str, text: str, error=False):
    print(("❌ " if error else "✅ ") + text)
    try:
        import tkinter as tk
        from tkinter import messagebox

        root = tk.Tk()
        root.withdraw()
        (messagebox.showerror if error else messagebox.showinfo)(title, text)
        root.destroy()
    except Exception:
        pass


def main():
    ap = argparse.ArgumentParser(description="สรุปยอดขายจากไฟล์ Excel/CSV")
    ap.add_argument("file", nargs="?", help="ไฟล์ยอดขาย (.xlsx/.xls/.csv)")
    ap.add_argument("-o", "--output", help="ชื่อไฟล์ผลลัพธ์ (ค่าเริ่มต้น: <ชื่อเดิม>_สรุป.xlsx)")
    ap.add_argument("--top", type=int, default=10, help="จำนวนสินค้าขายดีในกราฟ (ค่าเริ่มต้น 10)")
    args = ap.parse_args()

    src = args.file or pick_file_dialog()
    if not src:
        print("ไม่ได้เลือกไฟล์")
        sys.exit(1)
    src = Path(src)
    out = Path(args.output) if args.output else src.with_name(f"{src.stem}_สรุป.xlsx")

    try:
        df = clean(read_file(src))
        if df.empty:
            raise ValueError("ไม่มีข้อมูลที่ใช้ได้ในไฟล์")
        overview, products, daily = summarize(df, args.top)
        save_excel(overview, products, daily, args.top, out)
    except PermissionError:
        show_message("SalesSummary", f"บันทึกไฟล์ไม่ได้ กรุณาปิดไฟล์ {out.name} ก่อน แล้วลองใหม่", error=True)
        sys.exit(1)
    except Exception as e:
        show_message("SalesSummary", f"เกิดข้อผิดพลาด:\n{e}", error=True)
        sys.exit(1)

    dropped = df.attrs.get("dropped", 0)
    note = f"\n(ข้ามแถวที่ข้อมูลไม่ครบ {dropped} แถว)" if dropped else ""
    show_message("SalesSummary", f"สรุปเสร็จแล้ว\nบันทึกที่: {out}{note}")


if __name__ == "__main__":
    main()

       